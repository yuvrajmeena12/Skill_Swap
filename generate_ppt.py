import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
# Use a blank slide layout
blank_slide_layout = prs.slide_layouts[6]

def add_header_footer(slide):
    # Add SKIT Logo placeholder
    # left, top, width, height
    # We will just add a text box for the logo to avoid needing an actual image file
    logo_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.2), Inches(1.2), Inches(1.2))
    logo_box.fill.solid()
    logo_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    logo_box.line.color.rgb = RGBColor(0, 0, 0)
    text_frame = logo_box.text_frame
    text_frame.text = "SKIT Logo"
    text_frame.paragraphs[0].font.size = Pt(12)
    text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)

    # Add red footer bar
    footer_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.2), Inches(10), Inches(0.3))
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = RGBColor(192, 0, 0) # Red
    footer_bar.line.fill.background()
    
    # Add footer text "SKIT, JAIPUR"
    tx_left = slide.shapes.add_textbox(Inches(0.1), Inches(7.15), Inches(2), Inches(0.3))
    tf_left = tx_left.text_frame
    p_left = tf_left.paragraphs[0]
    p_left.text = "SKIT, JAIPUR"
    p_left.font.size = Pt(12)
    p_left.font.color.rgb = RGBColor(255, 255, 255)
    p_left.font.bold = True

    # Add footer text "www.skit.ac.in"
    tx_center = slide.shapes.add_textbox(Inches(4), Inches(7.15), Inches(2), Inches(0.3))
    tf_center = tx_center.text_frame
    p_center = tf_center.paragraphs[0]
    p_center.text = "www.skit.ac.in"
    p_center.font.size = Pt(12)
    p_center.font.color.rgb = RGBColor(255, 255, 255)
    p_center.font.bold = True
    p_center.alignment = PP_ALIGN.CENTER

def add_title(slide, title_text):
    # Add the red banner title as shown in the images
    title_banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(0.2), Inches(8.3), Inches(0.8))
    title_banner.fill.solid()
    title_banner.fill.fore_color.rgb = RGBColor(160, 40, 40)
    title_banner.line.color.rgb = RGBColor(0, 0, 0)
    
    tf = title_banner.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

# Slide 1: Title Slide (Topic: Skill Swap)
slide1 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide1)

# Swami Keshvanand Institute of Technology...
tx = slide1.shapes.add_textbox(Inches(1.5), Inches(0.2), Inches(8), Inches(1))
p = tx.text_frame.paragraphs[0]
p.text = "Swami Keshvanand Institute of Technology, Management & Gramothan, Jaipur"
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(192, 0, 0)
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

p2 = tx.text_frame.add_paragraph()
p2.text = "(An Autonomous Institute Affiliated to Rajasthan Technical University, Kota)"
p2.font.size = Pt(16)
p2.font.color.rgb = RGBColor(192, 0, 0)
p2.font.bold = True
p2.alignment = PP_ALIGN.CENTER

tx2 = slide1.shapes.add_textbox(Inches(1.5), Inches(2), Inches(7), Inches(1))
p3 = tx2.text_frame.paragraphs[0]
p3.text = "Industrial Training\nPresentation"
p3.font.size = Pt(28)
p3.font.bold = True
p3.alignment = PP_ALIGN.CENTER

tx3 = slide1.shapes.add_textbox(Inches(1), Inches(3.5), Inches(8), Inches(1))
p4 = tx3.text_frame.paragraphs[0]
p4.text = "TOPIC: Skill Swap"
p4.font.size = Pt(40)
p4.font.bold = True
p4.font.underline = True
p4.alignment = PP_ALIGN.CENTER

tx4 = slide1.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(4), Inches(1.5))
tf4 = tx4.text_frame
tf4.text = "Submitted To:\nName:- Mr. Vikash Dhankar\nAssistant Professor"
for paragraph in tf4.paragraphs:
    paragraph.font.size = Pt(16)
    paragraph.font.bold = True

tx5 = slide1.shapes.add_textbox(Inches(5.5), Inches(5.5), Inches(4), Inches(1.5))
tf5 = tx5.text_frame
tf5.text = "Submitted By:\nName:- Yuvraj Meena (24ESKCS780)"
for paragraph in tf5.paragraphs:
    paragraph.font.size = Pt(16)
    paragraph.font.bold = True

# Slide 2: Index
slide2 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide2)
add_title(slide2, "Index")

tx = slide2.shapes.add_textbox(Inches(2), Inches(1.5), Inches(6), Inches(5))
tf = tx.text_frame
items = [
    "1. Company Introduction",
    "2. Certificate",
    "3. Technology Used",
    "4. Goal of Training",
    "5. Project Overview",
    "6. Working Project",
    "7. Future Scope",
    "8. Q&A"
]
for i, item in enumerate(items):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = item
    p.font.size = Pt(28)

# Slide 3: Company Introduction (Leave blank)
slide3 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide3)
add_title(slide3, "Company Introduction")
tx = slide3.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(1))
tx.text_frame.text = "[Insert Company Introduction Here]"

# Slide 4: Certificate (Leave blank)
slide4 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide4)
add_title(slide4, "Certificate")
tx = slide4.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(1))
tx.text_frame.text = "[Insert Certificate Here]"

# Slide 5: Technology Used
slide5 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide5)
add_title(slide5, "Tech Stack")

tx = slide5.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(4))
tf = tx.text_frame
tf.text = "Technologies used in the project:"
p = tf.paragraphs[0]
p.font.size = Pt(24)

techs = [
    "Frontend: React.js, Tailwind CSS",
    "Backend: Node.js, Express.js",
    "Database: MongoDB",
    "Authentication: JWT (JSON Web Tokens)",
    "API Requests: Axios",
    "Other Tools: Vite, Cloudinary"
]
for tech in techs:
    p = tf.add_paragraph()
    p.text = "• " + tech
    p.font.size = Pt(20)

# Slide 6: Goal of Training
slide6 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide6)
add_title(slide6, "Goal of Training")

tx = slide6.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(5))
tf = tx.text_frame
goals = [
    "Learn practical Full-Stack Web Development",
    "Understand React frontend development",
    "Develop Node.js & Express backend APIs",
    "Work with MongoDB and Cloudinary",
    "Gain experience in building a complete real-world project"
]
for i, goal in enumerate(goals):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = "• " + goal
    p.font.size = Pt(24)

# Slide 7: Project Overview
slide7 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide7)
add_title(slide7, "Project Overview")

tx = slide7.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(5))
tf = tx.text_frame
p = tf.paragraphs[0]
p.text = "Skill Swap is a platform that allows users to exchange skills with each other. It provides a platform for individuals to offer their expertise and learn new skills from others without any monetary exchange."
p.font.size = Pt(24)

# Slide 8: Working Project (Screenshots)
slide8 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide8)
add_title(slide8, "Working Project")
tx = slide8.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(1))
tx.text_frame.text = "[Insert Project Screenshots Here]"

# Slide 9: Future Scope
slide9 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide9)
add_title(slide9, "Future Scope")

tx = slide9.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(5))
tf = tx.text_frame
scopes = [
    "Implement real-time chat between users",
    "Add user reviews and ratings system",
    "Integrate video calling for remote skill sharing",
    "Develop a mobile application for better accessibility"
]
for i, scope in enumerate(scopes):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = "• " + scope
    p.font.size = Pt(24)

# Slide 10: Q&A
slide10 = prs.slides.add_slide(blank_slide_layout)
add_header_footer(slide10)
add_title(slide10, "Q & A")
tx = slide10.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(1))
p = tx.text_frame.paragraphs[0]
p.text = "Any Questions?"
p.font.size = Pt(40)
p.alignment = PP_ALIGN.CENTER

prs.save("Skill_Swap_Presentation.pptx")
print("Presentation generated successfully!")
